#!/usr/bin/env python3
"""
TBBS-JIG スプレッドシートの一括セットアップ（Excel 経由）

Usage:
    python3 scripts/build-template-xlsx.py [input.xlsx] [output.xlsx]

  - input.xlsx：現状の Google Sheet を「ファイル → ダウンロード → .xlsx」した結果
  - output.xlsx：このスクリプトが整えた xlsx。事務局が Google Sheets に
    「ファイル → インポート → 既存のシートを置き換え」でアップロード。

何をするか:
  1. Issue主檔 シートに、不足列（Confluence URL / 狀態 / 事務局備註 / 更新日）を追加
  2. Task一覧 シートを新規作成（または再作成）し、ヘッダ 8 列を整える
  3. 両シートの狀態列にプルダウン（未開始/進行中/完成）と3色条件付き書式
  4. 両シートの更新日列に鮮度色変化（>7日黄、>14日赤）
  5. 列幅・ヘッダ行 freeze

データ検証と条件付き書式は xlsx ネイティブの形式で保存するため、
Google Sheets にインポートしても継承される（プルダウンも色も動く）。
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

INPUT  = Path(sys.argv[1] if len(sys.argv) > 1 else 'Issue主檔.xlsx')
OUTPUT = Path(sys.argv[2] if len(sys.argv) > 2 else 'JIG-template.xlsx')

NEW_ISSUE_COLS = ['Confluence URL', '狀態', '事務局備註', '更新日']
STATUS_VALUES  = ['未開始', '進行中', '完成']
STATUS_COLORS  = {'未開始': 'EEEAE0', '進行中': 'DCEBFB', '完成': 'D4F2DD'}
TASK_HEADERS   = ['編號', 'Task', '親編號', '戰略負責人', '狀態', '事務局備註', '更新日', 'Confluence URL']
TASK_WIDTHS    = [12, 38, 12, 12, 10, 28, 12, 30]
HDR_FILL = PatternFill(start_color='F7F4EC', end_color='F7F4EC', fill_type='solid')
HDR_FONT = Font(bold=True)

def apply_status_dropdown(ws, col_letter, row_count=500):
    dv = DataValidation(
        type='list',
        formula1='"{}"'.format(','.join(STATUS_VALUES)),
        allow_blank=True,
        showErrorMessage=True,
    )
    dv.error = '未開始 / 進行中 / 完成 のいずれかを選択してください'
    dv.errorTitle = '無効な値'
    ws.add_data_validation(dv)
    dv.add('{c}2:{c}{n}'.format(c=col_letter, n=row_count))

def apply_status_colors(ws, col_letter, row_count=500):
    rng = '{c}2:{c}{n}'.format(c=col_letter, n=row_count)
    for v, color in STATUS_COLORS.items():
        rule = CellIsRule(
            operator='equal',
            formula=['"{}"'.format(v)],
            fill=PatternFill(start_color=color, end_color=color, fill_type='solid'),
        )
        ws.conditional_formatting.add(rng, rule)

def apply_freshness_colors(ws, col_letter, row_count=500):
    rng = '{c}2:{c}{n}'.format(c=col_letter, n=row_count)
    # 7日超〜14日：黄
    warn = FormulaRule(
        formula=[
            'AND(${c}2<>"", TODAY()-${c}2>7, TODAY()-${c}2<=14)'.format(c=col_letter)
        ],
        fill=PatternFill(start_color='FFF3CC', end_color='FFF3CC', fill_type='solid'),
        font=Font(color='7A5A00'),
    )
    # 14日超：赤
    stale = FormulaRule(
        formula=[
            'AND(${c}2<>"", TODAY()-${c}2>14)'.format(c=col_letter)
        ],
        fill=PatternFill(start_color='FCD7D7', end_color='FCD7D7', fill_type='solid'),
        font=Font(color='8A0000'),
    )
    ws.conditional_formatting.add(rng, warn)
    ws.conditional_formatting.add(rng, stale)

def header_idx(ws, name):
    for cell in ws[1]:
        if cell.value == name:
            return cell.column
    return None

# ---- 処理開始 ----
if not INPUT.exists():
    print(f'入力ファイルが見つかりません: {INPUT}', file=sys.stderr)
    sys.exit(1)

wb = openpyxl.load_workbook(INPUT)

# 1) Issue主檔 に不足列を追加
if 'Issue主檔' not in wb.sheetnames:
    print('「Issue主檔」シートがありません', file=sys.stderr)
    sys.exit(1)
issue_ws = wb['Issue主檔']
existing = [c.value for c in issue_ws[1]]
for name in NEW_ISSUE_COLS:
    if name not in existing:
        col = issue_ws.max_column + 1
        cell = issue_ws.cell(row=1, column=col, value=name)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        existing.append(name)
        print(f'  + Issue主檔 に列「{name}」を追加')

# 2) Issue主檔 のフォーマット
st_idx  = header_idx(issue_ws, '狀態')
upd_idx = header_idx(issue_ws, '更新日')
if st_idx:
    apply_status_dropdown(issue_ws, get_column_letter(st_idx), row_count=500)
    apply_status_colors  (issue_ws, get_column_letter(st_idx), row_count=500)
if upd_idx:
    apply_freshness_colors(issue_ws, get_column_letter(upd_idx), row_count=500)
issue_ws.freeze_panes = 'A2'

# 3) Task一覧 シートを作成（既にあれば消して作り直し）
if 'Task一覧' in wb.sheetnames:
    del wb['Task一覧']
task_ws = wb.create_sheet('Task一覧')

for i, h in enumerate(TASK_HEADERS, start=1):
    cell = task_ws.cell(row=1, column=i, value=h)
    cell.font = HDR_FONT
    cell.fill = HDR_FILL

for i, w in enumerate(TASK_WIDTHS, start=1):
    task_ws.column_dimensions[get_column_letter(i)].width = w

apply_status_dropdown(task_ws, 'E', row_count=500)
apply_status_colors  (task_ws, 'E', row_count=500)
apply_freshness_colors(task_ws, 'G', row_count=500)
task_ws.freeze_panes = 'A2'
print('  + Task一覧 シートを作成（8列）')

wb.save(OUTPUT)
print(f'\n✅ 出力完了: {OUTPUT}')
print('   Google Sheets で「ファイル → インポート → 既存のシートを置き換え」でアップロード')
